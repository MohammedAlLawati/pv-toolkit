"""Build a letter (.docx) on the company letterhead from a template and the form values."""
from __future__ import annotations

import io
import re
from contextvars import ContextVar
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ASSETS = Path(__file__).resolve().parent / "assets"
DEFAULT_TOP = ASSETS / "htc_letterhead_top.png"
DEFAULT_BOTTOM = ASSETS / "htc_letterhead_bottom.png"

# placeholder -> label, shown on the template editor
PLACEHOLDERS = {
    "Ref": "Letter reference",
    "Date": "Letter date",
    "ClientName": "Client / company / homeowner name",
    "CRNo": "Company C.R. number",
    "SignatoryName": "Authorised signatory's name",
    "SignatoryID": "Authorised signatory's ID",
    "CivilID": "Homeowner's civil ID",
    "Location": "Premises location",
    "Purpose": "What the letter is about",
    "MeterList": "Meter numbers, comma separated",
    "AccountList": "Account numbers, comma separated",
    "CompanyName": "HTC company name (in the letter's language)",
}
BLOCKS = ("details", "points", "enclosures", "signatures", "htc_signature", "client_signature")

NAVY = RGBColor(0x1F, 0x3A, 0x68)
GREY = RGBColor(0x59, 0x59, 0x59)
FONT = "Calibri"
FONT_AR = "Noto Naskh Arabic"    # shipped in assets/fonts so PDFs render the same everywhere
AR_SIZE_BOOST = 1.5               # Arabic reads smaller than Latin at the same point size
BLANK = "____________________"   # an empty input, printed so it can be filled by hand

# fixed wording around the template text, per letter language
LABELS = {
    "en": {
        "ref": "Ref:", "date": "Date:", "subject": "Subject:", "company": "Hussam Technology Company L.L.C.",
        "company_name": "Company name", "cr": "C.R. No.", "signatory": "Authorised signatory",
        "signatory_id": "Signatory ID No.", "homeowner": "Homeowner", "civil_id": "Civil ID No.",
        "location": "Premises location", "no": "No.", "meter": "Meter No.", "account": "Account No.",
        "meter_opt": "Meter No.", "place": "Location / description", "enclosures": "Enclosures:",
        "faithfully": "Yours faithfully,", "for": "For", "name": "Name", "designation": "Designation",
        "signature": "Signature", "stamp": "Company stamp", "id": "ID No.",
        "client_box": "CLIENT CONFIRMATION", "owner_box": "HOMEOWNER CONFIRMATION",
    },
    "ar": {
        "ref": "الرقم المرجعي:", "date": "التاريخ:", "subject": "الموضوع:",
        "company": "شركة حسام للتقنية ش.م.م", "company_name": "اسم الشركة",
        "cr": "رقم السجل التجاري", "signatory": "المفوّض بالتوقيع",
        "signatory_id": "رقم بطاقة المفوّض", "homeowner": "اسم المالك",
        "civil_id": "الرقم المدني", "location": "موقع المبنى", "no": "م", "meter": "رقم العداد",
        "account": "رقم الحساب", "meter_opt": "رقم العداد", "place": "الموقع / الوصف",
        "enclosures": "المرفقات:", "faithfully": "وتفضلوا بقبول فائق الاحترام والتقدير،",
        "for": "عن", "name": "الاسم", "designation": "المسمى الوظيفي", "signature": "التوقيع",
        "stamp": "ختم الشركة", "id": "رقم البطاقة", "client_box": "إقرار العميل",
        "owner_box": "إقرار المالك",
    },
}

_RTL: ContextVar[bool] = ContextVar("letter_rtl", default=False)


@dataclass
class SupplyPoint:
    meter_no: str = ""
    account_no: str = ""
    location: str = ""


@dataclass
class LetterData:
    audience: str                       # commercial | residential
    recipient: str
    subject: str
    body: str
    confirmation: str = ""
    enclosures: list[str] = field(default_factory=list)
    values: dict[str, str] = field(default_factory=dict)
    points: list[SupplyPoint] = field(default_factory=list)
    language: str = "en"                # en | ar (right to left)

    @property
    def labels(self) -> dict[str, str]:
        return LABELS.get(self.language, LABELS["en"])


def complete_values(d: LetterData) -> dict[str, str]:
    v = {k: "" for k in PLACEHOLDERS}
    v["CompanyName"] = d.labels["company"]
    v.update({k: (val or "") for k, val in d.values.items()})
    v["MeterList"] = ", ".join(p.meter_no for p in d.points if p.meter_no)
    v["AccountList"] = ", ".join(p.account_no for p in d.points if p.account_no)
    return v


def fill(text: str, values: dict[str, str], missing: set[str] | None = None) -> str:
    def sub(m: re.Match) -> str:
        key = m.group(1).strip()
        val = values.get(key, "")
        if not val and missing is not None:
            missing.add(key)
        return val or BLANK
    return re.sub(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}", sub, text)


# ---- low-level docx helpers --------------------------------------------------------------

def _runs(par, text: str, size: float = 10.5, color: RGBColor | None = None, bold: bool = False):
    """Add text to a paragraph, honouring **bold** markers."""
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if not chunk:
            continue
        r = par.add_run(chunk)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold or (i % 2 == 1)
        if color is not None:
            r.font.color.rgb = color
        if _RTL.get():
            r_pr = r._r.get_or_add_rPr()
            r_pr.rFonts.set(qn("w:cs"), FONT_AR)
            for tag, val in (("w:bCs", None if r.bold else "0"),
                             ("w:szCs", str(int((size + AR_SIZE_BOOST) * 2)))):
                el = OxmlElement(tag)
                if val is not None:
                    el.set(qn("w:val"), val)
                r_pr.append(el)
            r_pr.append(OxmlElement("w:rtl"))
    if _RTL.get():
        _bidi(par)
        # Arabic fonts have tall ascenders; a fixed line height keeps the page compact
        par.paragraph_format.line_spacing = Pt((size + AR_SIZE_BOOST) * 1.5)


def _bidi(par) -> None:
    """Mark a paragraph right to left (once)."""
    p_pr = par._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.insert(0, OxmlElement("w:bidi"))


def _para(doc, text: str = "", *, size: float = 10.5, bold=False, color=None, after: float = 6,
          align=None, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    if text:
        _runs(p, text, size=size, color=color, bold=bold)
    elif _RTL.get():
        _bidi(p)
    return p


def _gap(target, pt: float):
    """An empty paragraph exactly pt points tall."""
    p = target.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = Pt(pt)
    return p


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _borders(table, color: str = "A6A6A6", size: int = 4, inside: bool = True) -> None:
    tbl_pr = table._tbl.tblPr
    if _RTL.get() and tbl_pr.find(qn("w:bidiVisual")) is None:
        tbl_pr.append(OxmlElement("w:bidiVisual"))  # columns run right to left
    b = OxmlElement("w:tblBorders")
    edges = ("top", "left", "bottom", "right") + (("insideH", "insideV") if inside else ())
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge in edges:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tbl_pr.append(b)


def _widths(table, widths: list[float]) -> None:
    """Fix column widths so Word and LibreOffice both honour them."""
    table.autofit = False
    tbl = table._tbl
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    tbl_pr = tbl.tblPr
    tw = tbl_pr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tbl_pr.append(tw)
    tw.set(qn("w:w"), str(int(Cm(sum(widths)).twips)))
    tw.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = Cm(w)


def _cell_text(cell, text: str, *, bold=False, size: float = 10, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    if align is not None:
        p.alignment = align
    _runs(p, text, size=size, bold=bold, color=color)
    if _RTL.get():
        _bidi(p)


def _table(doc, header: list[str] | None, rows: list[list[str]], widths: list[float],
           label_col: bool = False):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    if header:
        cells = t.add_row().cells
        for c, h in zip(cells, header):
            _cell_text(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            _shade(c, "1F3A68")
    for row in rows:
        cells = t.add_row().cells
        for j, (c, val) in enumerate(zip(cells, row)):
            is_label = label_col and j % 2 == 0
            _cell_text(c, val, bold=is_label)
            if is_label:
                _shade(c, "E8EDF5")
    _widths(t, widths)
    _gap(doc, 9)
    return t


# ---- blocks ------------------------------------------------------------------------------

def _block_details(doc, d: LetterData, v: dict[str, str]) -> None:
    L = d.labels
    if d.audience == "commercial":
        rows = [[L["company_name"], v.get("ClientName", ""), L["cr"], v.get("CRNo", "")],
                [L["signatory"], v.get("SignatoryName", ""), L["signatory_id"], v.get("SignatoryID", "")]]
    else:
        rows = [[L["homeowner"], v.get("ClientName", ""), L["civil_id"], v.get("CivilID", "")]]
    rows.append([L["location"], v.get("Location", ""), "", ""])
    widths = [3.4, 5.4, 3.6, 4.2] if d.language == "ar" else [3.8, 5.8, 3.0, 4.0]
    t = _table(doc, None, rows, widths, label_col=True)
    last = t.rows[-1].cells
    last[1].merge(last[3])
    _cell_text(last[1], v.get("Location", ""))


def _block_points(doc, d: LetterData, v: dict[str, str]) -> None:
    """Electricity accounts: account number and location / description."""
    L = d.labels
    pts = [p for p in d.points if p.account_no or p.location]
    if not pts:  # nothing entered yet: empty rows to fill by hand
        pts = [SupplyPoint() for _ in range(3 if d.audience == "commercial" else 2)]
    widths = [1.2, 6.4, 16.6 - 7.6]
    rows = [[str(n), p.account_no, p.location] for n, p in enumerate(pts, 1)]
    t = _table(doc, [L["no"], L["account"], L["place"]], rows, widths)
    for row in t.rows:
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for c in row.cells[1:]:
            c.paragraphs[0].paragraph_format.line_spacing = Pt(15)  # room to write by hand


def _enclosure_list(target, d: LetterData, v: dict[str, str], size: float = 10) -> None:
    items = [e for e in d.enclosures if e.strip()]
    if not items:
        return
    _para(target, d.labels["enclosures"], bold=True, after=1, size=size)
    for n, e in enumerate(items, 1):
        p = _para(target, f"{n}.  {fill(e, v)}", after=0, size=size)
        p.paragraph_format.left_indent = Cm(0.3)


def _block_enclosures(doc, d: LetterData, v: dict[str, str]) -> None:
    if any(e.strip() for e in d.enclosures):
        _enclosure_list(doc, d, v)
        _gap(doc, 8)


def _htc_signoff(target, d: LetterData, v: dict[str, str]) -> None:
    """HTC's sign-off, into the document or a table cell. Who signs is left blank."""
    L = d.labels
    _para(target, L["faithfully"], after=2)
    _para(target, f"{L['for']} **{v.get('CompanyName')}**", after=10)
    for label in ("name", "designation", "signature"):
        _para(target, f"**{L[label]}:**  {BLANK}", after=6, size=9.5)
    _para(target, f"**{L['stamp']}:**", after=0, size=9.5)


def _box_borders(cell, color: str = "1F3A68", size: int = 8) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        b.append(el)
    tc_pr.append(b)


def _keep_row(table) -> None:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:cantSplit")
        tr_pr.append(el)


def _confirmation(cell, d: LetterData, v: dict[str, str]) -> None:
    """Fill a table cell with the client's confirmation and signature lines."""
    L = d.labels
    who = L["client_box"] if d.audience == "commercial" else L["owner_box"]
    _box_borders(cell)
    _cell_text(cell, who, bold=True, size=9.5, color=NAVY)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    _runs(p, fill(d.confirmation, v), size=9.5)
    if d.audience == "commercial":
        lines = [(L["signatory"], v.get("SignatoryName")), (L["id"], v.get("SignatoryID")),
                 (L["signature"], ""), (L["date"].rstrip(":"), ""), (L["stamp"], None)]
    else:
        lines = [(L["homeowner"], v.get("ClientName")), (L["civil_id"], v.get("CivilID")),
                 (L["signature"], ""), (L["date"].rstrip(":"), "")]
    for label, val in lines:
        q = cell.add_paragraph()
        q.paragraph_format.space_after = Pt(5 if val is not None else 22)
        _runs(q, f"{label}:  ", size=9.5, bold=True)
        if val is not None:
            _runs(q, val or BLANK, size=9.5)


def _block_htc_signature(doc, d: LetterData, v: dict[str, str]) -> None:
    _htc_signoff(doc, d, v)
    _gap(doc, 8)


def _block_client_signature(doc, d: LetterData, v: dict[str, str]) -> None:
    if not d.confirmation.strip():
        return
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t, color="FFFFFF", size=0, inside=False)
    _widths(t, [16.6])
    _confirmation(t.rows[0].cells[0], d, v)
    _keep_row(t)


def _block_signatures(doc, d: LetterData, v: dict[str, str]) -> None:
    """HTC sign-off on the left, the client's confirmation box on the right."""
    if not d.confirmation.strip():
        return _block_htc_signature(doc, d, v)
    t = doc.add_table(rows=1, cols=3)
    t.autofit = False
    _borders(t, color="FFFFFF", size=0, inside=False)
    left, gap, right = t.rows[0].cells
    _widths(t, [7.4, 0.5, 8.7])
    left.paragraphs[0].paragraph_format.space_after = Pt(0)
    _htc_signoff(left, d, v)
    if "[[enclosures]]" not in d.body:  # enclosures sit under the sign-off unless placed elsewhere
        _gap(left, 14)
        _enclosure_list(left, d, v, size=9)
    _confirmation(right, d, v)
    _keep_row(t)


BLOCK_FUNCS = {
    "details": _block_details,
    "points": _block_points,
    "enclosures": _block_enclosures,
    "htc_signature": _block_htc_signature,
    "client_signature": _block_client_signature,
    "signatures": _block_signatures,
}


# ---- page setup --------------------------------------------------------------------------

def trim_blank_rows(data: bytes, pad: int = 4) -> bytes:
    """Cut empty (white or transparent) rows off the top and bottom of a letterhead image."""
    from PIL import Image, ImageChops
    im = Image.open(io.BytesIO(data))
    im.load()
    rgba = im.convert("RGBA")
    white = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
    flat = Image.alpha_composite(white, rgba).convert("L")
    ink = ImageChops.invert(flat).point(lambda x: 255 if x > 12 else 0)
    box = ink.getbbox()
    if not box:
        return data
    top, bottom = max(0, box[1] - pad), min(im.height, box[3] + pad)
    if top == 0 and bottom == im.height:
        return data
    out = io.BytesIO()
    rgba.crop((0, top, im.width, bottom)).save(out, "PNG")
    return out.getvalue()


def _letterhead(doc, top: bytes | None, bottom: bytes | None) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(3.9)
    sec.bottom_margin = Cm(2.9)
    sec.header_distance = Cm(0.4)
    sec.footer_distance = Cm(0.3)
    bleed = Cm(2.0)  # images run almost edge to edge
    for part, img in ((sec.header, top), (sec.footer, bottom)):
        part.is_linked_to_previous = False
        p = part.paragraphs[0]
        p.paragraph_format.left_indent = -bleed
        p.paragraph_format.right_indent = -bleed
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if img:
            img = trim_blank_rows(img)
            p.add_run().add_picture(io.BytesIO(img), width=Cm(21.0 - 4.4 + 2 * 2.0) - Cm(0.6))


def build_docx(d: LetterData, top: bytes | None = None, bottom: bytes | None = None
               ) -> tuple[bytes, set[str]]:
    """Return the letter as .docx bytes and the placeholders that had no value."""
    token = _RTL.set(d.language == "ar")
    try:
        return _build(d, top, bottom)
    finally:
        _RTL.reset(token)


def _build(d: LetterData, top: bytes | None, bottom: bytes | None) -> tuple[bytes, set[str]]:
    if top is None and DEFAULT_TOP.exists():
        top = DEFAULT_TOP.read_bytes()
    if bottom is None and DEFAULT_BOTTOM.exists():
        bottom = DEFAULT_BOTTOM.read_bytes()
    L = d.labels
    rtl = _RTL.get()
    v = complete_values(d)
    missing: set[str] = set()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.element.rPr.rFonts.set(qn("w:cs"), FONT_AR if rtl else FONT)
    _letterhead(doc, top, bottom)

    # reference (only when the template numbers its letters) and date on one line
    end = WD_ALIGN_PARAGRAPH.LEFT if rtl else WD_ALIGN_PARAGRAPH.RIGHT
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    _borders(t, color="FFFFFF", size=0, inside=False)
    left, right = t.rows[0].cells
    _widths(t, [8.3, 8.3])
    if v.get("Ref"):
        _cell_text(left, f"**{L['ref']}** {v['Ref']}", size=10.5)
    elif rtl:
        _bidi(left.paragraphs[0])
    _cell_text(right, f"**{L['date']}** {fill('{{Date}}', v, missing)}", size=10.5, align=end)
    _gap(doc, 6)

    lines = [ln for ln in fill(d.recipient, v, missing).splitlines() if ln.strip()]
    for n, ln in enumerate(lines):
        _para(doc, ln, after=0, bold=(n == len(lines) - 2 and len(lines) > 1))
    _gap(doc, 10)

    subj = _para(doc, after=8)
    _runs(subj, f"{L['subject']} ", bold=True)
    _runs(subj, fill(d.subject, v, missing), bold=True)
    for r in subj.runs:
        r.underline = True

    for chunk in re.split(r"\n\s*\n", d.body.strip()):
        chunk = chunk.strip()
        m = re.fullmatch(r"\[\[\s*([a-z_]+)\s*\]\]", chunk)
        if m:
            fn = BLOCK_FUNCS.get(m.group(1))
            if fn:
                fn(doc, d, v)
            continue
        text = fill(chunk, v, missing)
        if all(ln.lstrip().startswith("- ") for ln in text.splitlines()):
            for ln in text.splitlines():
                p = _para(doc, "•  " + ln.lstrip()[2:], after=2)
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.4)
            _gap(doc, 6)
        else:
            _para(doc, " ".join(text.splitlines()), after=7, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # placeholders used only inside blocks
    if "[[details]]" in d.body:
        keys = ("ClientName", "CRNo", "SignatoryName", "SignatoryID") if d.audience == "commercial" \
            else ("ClientName", "CivilID")
        missing.update(k for k in keys if not v.get(k))
    if "[[points]]" in d.body and not (v["MeterList"] or v["AccountList"]):
        missing.add("MeterList" if d.audience == "commercial" else "AccountList")
    # shrink the trailing paragraph Word needs after a closing table, so it can't spill a page
    last = doc.paragraphs[-1]
    if not last.text:
        last.paragraph_format.space_after = Pt(0)
        last.paragraph_format.line_spacing = Pt(1)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), missing
